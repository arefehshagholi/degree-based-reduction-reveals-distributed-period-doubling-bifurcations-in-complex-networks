"""Regenerate network_ranges.csv from the bundled critical-point workbook.

Requires Python 3 and openpyxl.
Run from this directory with: python build_range_catalog.py
"""
from pathlib import Path
import csv
import math
import openpyxl

HERE = Path(__file__).resolve().parent


def network_filename(label):
    if label == "Complete":
        return "completenet.mat"
    if label == "Star":
        return "starnet.mat"
    for prefix, stem in [
        ("Regular", "regularnet"),
        ("WS", "WattsStrogatz"),
        ("NW", "NewmanWatts"),
        ("SF", "BAnet"),
    ]:
        if label.startswith(prefix):
            return stem + label[len(prefix):] + ".mat"
    raise ValueError(f"Unknown network label: {label}")


def build():
    workbook = openpyxl.load_workbook(
        HERE / "manually_created_critical_points.xlsx", data_only=True
    )
    sheet_map = {
        "Rossler system": "rossler",
        "mChialvo system": "mchialvo",
        "Henon system": "henon",
    }

    rows = []
    for sheet in workbook.worksheets:
        title = sheet.title.strip()
        if title not in sheet_map:
            continue
        system = sheet_map[title]
        groups = {}

        for row in range(2, sheet.max_row + 1):
            label = sheet.cell(row, 1).value
            if not label:
                continue
            label = str(label).strip()
            group = groups.setdefault(label, [])

            cp_columns = [4, 6] if system == "henon" else [4, 6, 8]
            for column in cp_columns:
                for col in (column, column + 1):
                    value = sheet.cell(row, col).value
                    if isinstance(value, (int, float)) and math.isfinite(value) and value > 0:
                        group.append(float(value))

        if len(groups) != 22:
            raise ValueError(f"{system}: expected 22 networks, found {len(groups)}")

        for label, values in groups.items():
            if not values:
                raise ValueError(f"No positive finite critical points for {system}/{label}")
            filename = network_filename(label)
            network_path = HERE / "networks" / filename
            if not network_path.is_file():
                raise FileNotFoundError(network_path)
            rows.append([
                system,
                label,
                filename,
                0,
                format(1.10 * max(values), ".15g"),
            ])

    with (HERE / "network_ranges.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow([
            "system",
            "network_label",
            "network_file",
            "suggested_min",
            "suggested_max",
        ])
        writer.writerows(rows)

    print("Saved coupling-range catalog to network_ranges.csv")


if __name__ == "__main__":
    build()
